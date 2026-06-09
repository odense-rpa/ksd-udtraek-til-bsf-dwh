from datetime import date, datetime
from typing import Dict, List
from openpyxl import load_workbook


excel_mappings: List[Dict[str, str]] = []


def _format_excel_value(raw_value: object) -> str:
    if raw_value is None:
        return ""
    if isinstance(raw_value, datetime):
        return raw_value.strftime("%d-%m-%Y")
    if isinstance(raw_value, date):
        return raw_value.strftime("%d-%m-%Y")
    return str(raw_value).strip()


def get_excel_mapping() -> List[Dict[str, str]]:
    """Henter POF-mapping fra Martin og Sofies regneark"""
    global excel_mappings
    if not excel_mappings:
        raise ValueError("excel-mapping er ikke indlæst, brug load_excel_mapping først")
    return excel_mappings


def load_excel_mapping(file_path: str):
    global excel_mappings
    try:
        workbook = load_workbook(file_path)
        worksheet = workbook.active

        if worksheet is None:
            raise ValueError("Worksheet could not be loaded")

        # Read column headers from first row.
        header_row = worksheet[1]
        headers: List[str] = []
        for cell in header_row:
            if cell.value and str(cell.value).strip():
                headers.append(str(cell.value).strip())

        if not headers:
            raise ValueError("No headers found in row 1")

        rows: List[Dict[str, str]] = []
        for row in worksheet.iter_rows(
            min_row=2,
            max_col=len(headers),
            values_only=True,
        ):
            row_data: Dict[str, str] = {}
            has_value = False

            for idx, header in enumerate(headers):
                raw_value = row[idx] if idx < len(row) else None
                value = _format_excel_value(raw_value)
                if value:
                    has_value = True
                row_data[header] = value

            if has_value:
                rows.append(row_data)

        excel_mappings = rows

    except Exception as e:
        raise RuntimeError(
            f"Failed to load mapping from Excel file '{file_path}': {str(e)}"
        ) from e
